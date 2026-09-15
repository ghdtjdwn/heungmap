"""문체부 연도별 지역축제 XLSX를 보수적으로 정규화하고 시계열 학습쌍을 만든다."""

from __future__ import annotations

import math
import re
import unicodedata
from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from zipfile import ZipFile

import numpy as np
import openpyxl
import pandas as pd


SUPPORTED_YEARS = tuple(range(2017, 2027))
SOURCE_URL = "https://www.mcst.go.kr/site/s_culture/festival/festivalList.jsp"
DOWNLOAD_URL = "https://www.mcst.go.kr/servlets/eduport/front/upload/UplDownloadFile"
FEATURE_VERSION = "mcst-attendance-recurring-v1"
PROVINCES = (
    "서울", "부산", "대구", "인천", "광주", "대전", "울산", "세종", "경기",
    "강원", "충북", "충남", "전북", "전남", "경북", "경남", "제주", "unknown",
)
FESTIVAL_TYPES = ("문화예술", "전통역사", "주민화합", "자연생태", "지역특산물", "기타", "unknown")
NUMERIC_FEATURES = (
    "log_prior_attendance", "log_budget", "log_budget_per_prior_visitor",
    "festival_age", "month_sin", "month_cos", "duration_days",
)
FEATURES = [
    *NUMERIC_FEATURES,
    *(f"province_{index}" for index in range(len(PROVINCES))),
    *(f"festival_type_{index}" for index in range(len(FESTIVAL_TYPES))),
]

_PROVINCE_ALIASES = {
    "서울특별시": "서울", "부산광역시": "부산", "대구광역시": "대구", "인천광역시": "인천",
    "광주광역시": "광주", "대전광역시": "대전", "울산광역시": "울산", "세종특별자치시": "세종",
    "경기도": "경기", "강원도": "강원", "강원특별자치도": "강원", "충청북도": "충북",
    "충청남도": "충남", "전라북도": "전북", "전북특별자치도": "전북", "전라남도": "전남",
    "경상북도": "경북", "경상남도": "경남", "제주특별자치도": "제주",
}


def text(value: Any) -> str:
    return " ".join(str(value or "").split())


def number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    cleaned = text(value).replace(",", "")
    if not cleaned or cleaned in {"-", "미집계", "미정", "없음", "해당없음", "추산불가"}:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", cleaned)
    return float(match.group()) if match else None


def canonical_title(value: Any) -> str:
    """회차·연도·구두점을 지우되 의미 단어가 같은 경우만 같은 축제로 본다."""
    value = unicodedata.normalize("NFKC", text(value)).lower()
    value = re.sub(r"(?:19|20)\d{2}", "", value)
    value = re.sub(r"제\s*\d+\s*회|\d+\s*회", "", value)
    value = re.sub(r"\d+", "", value)
    return "".join(re.findall(r"[가-힣a-z]+", value))


def normalize_province(value: Any) -> str:
    value = re.sub(r"^\d+\.\s*", "", text(value))
    return _PROVINCE_ALIASES.get(value, value if value in PROVINCES else "unknown")


def normalize_district(value: Any) -> str:
    value = text(value)
    return "" if value in {"-", "본청", "시자체", "도자체", "시 자체", "도 자체"} else value


def normalize_type(value: Any) -> str:
    value = re.sub(r"^\d+\.\s*", "", text(value)).replace("·", "").replace(" ", "")
    aliases = {"문화·예술": "문화예술", "생태자원": "자연생태", "산업,문화축제": "기타"}
    value = aliases.get(value, value)
    return value if value in FESTIVAL_TYPES else "기타" if value else "unknown"


def _year(value: Any) -> int | None:
    match = re.search(r"(?:19|20)\d{2}", text(value))
    return int(match.group()) if match else None


def _month(value: Any) -> int:
    match = re.search(r"(?<!\d)(1[0-2]|0?[1-9])\s*[./월]", text(value))
    return int(match.group(1)) if match else 6


def _duration(value: Any) -> int:
    match = re.search(r"(\d+)\s*일간", text(value))
    return max(1, min(90, int(match.group(1)))) if match else 1


def _measurement(value: Any) -> str:
    value = text(value)
    for key in ("계측", "추정", "미집계", "기타", "무응답"):
        if key in value:
            return key
    return "미제공"


def _record(*, plan_year: int, province: Any, district: Any, title: Any, festival_type: Any,
            budget: Any, first_year: Any, month: Any, duration: Any, attendance: Any,
            measurement: Any = None) -> dict[str, Any] | None:
    title = text(title)
    key = canonical_title(title)
    if not title or len(key) < 2:
        return None
    attendance = number(attendance)
    return {
        "plan_year": plan_year, "province": normalize_province(province),
        "district": normalize_district(district), "title": title, "title_key": key,
        "festival_type": normalize_type(festival_type), "budget_million_krw": number(budget),
        "first_year": _year(first_year), "month": int(month) if number(month) and 1 <= int(float(month)) <= 12 else 6,
        "duration_days": max(1, min(90, int(float(duration)))) if number(duration) else 1,
        "prior_attendance": attendance, "measurement": _measurement(measurement),
    }


def _old_records(year: int, workbook: openpyxl.Workbook) -> list[dict[str, Any]]:
    attendance_column = {2017: 21, 2018: 23, 2019: 23, 2020: 22, 2021: 22}[year]
    budget_column = 10 if year == 2017 else 9
    type_column = 15 if year == 2017 else 16
    first_year_column = 9 if year == 2017 else 8
    rows = []
    for sheet in workbook.worksheets:
        if sheet.title in {"총괄", "총괄표"}:
            continue
        for values in sheet.iter_rows(min_row=5, values_only=True):
            if len(values) < attendance_column or number(values[0]) is None:
                continue
            attendance = number(values[attendance_column - 1])
            if attendance is not None:
                attendance *= 1_000  # 2017~2021 파일은 천 명 단위다.
            record = _record(
                plan_year=year, province=values[1], district=values[2], title=values[3],
                festival_type=values[type_column - 1], budget=values[budget_column - 1],
                first_year=values[first_year_column - 1], month=_month(values[4]),
                duration=_duration(values[5] if len(values) > 5 else None), attendance=attendance,
            )
            if record:
                rows.append(record)
    return rows


def _new_records(year: int, workbook: openpyxl.Workbook) -> list[dict[str, Any]]:
    sheet = workbook["조사표"] if year >= 2025 else workbook["세부현황"]
    start_row = {2022: 8, 2023: 8, 2024: 7, 2025: 8, 2026: 9}[year]
    rows = []
    for values in sheet.iter_rows(min_row=start_row, values_only=True):
        if len(values) < 16 or number(values[1]) is None:
            continue
        if year in {2022, 2023}:
            domestic, foreign = number(values[14]), number(values[15])
            attendance = ((domestic or 0) + (foreign or 0)) if domestic is not None or foreign is not None else None
            budget, first_year, month, duration = values[13], values[9], _month(values[6]), _duration(values[6])
        elif year == 2024:
            attendance, budget, first_year, month, duration = values[14], values[10], values[9], _month(values[6]), _duration(values[6])
        elif year == 2025:
            attendance, budget, first_year, month, duration = values[26], values[22], values[21], values[12], values[17]
        else:
            attendance, budget, first_year, month, duration = values[26], values[21], values[20], values[12], values[17]
        record = _record(
            plan_year=year, province=values[2], district=values[3], title=values[4],
            festival_type=values[5], budget=budget, first_year=first_year, month=month,
            duration=duration, attendance=attendance,
            measurement=values[29] if year == 2026 and len(values) > 29 else None,
        )
        if record:
            rows.append(record)
    return rows


def load_archives(paths: Iterable[Path]) -> tuple[pd.DataFrame, dict[str, Any]]:
    records: list[dict[str, Any]] = []
    archive_rows: dict[str, int] = {}
    seen_years: set[int] = set()
    for path in sorted(map(Path, paths)):
        match = re.match(r"((?:19|20)\d{2})_festival\.zip$", path.name)
        if not match:
            raise ValueError(f"연도를 판별할 수 없는 문체부 파일입니다: {path.name}")
        year = int(match.group(1))
        if year not in SUPPORTED_YEARS or year in seen_years:
            raise ValueError("문체부 지원 연도가 아니거나 같은 연도가 중복됐습니다.")
        seen_years.add(year)
        with ZipFile(path) as archive:
            books = [name for name in archive.namelist() if name.lower().endswith(".xlsx")]
            if len(books) != 1:
                raise ValueError(f"XLSX가 정확히 하나가 아닙니다: {path.name}")
            workbook = openpyxl.load_workbook(BytesIO(archive.read(books[0])), read_only=True, data_only=True)
            annual = _old_records(year, workbook) if year <= 2021 else _new_records(year, workbook)
        if not annual:
            raise ValueError(f"축제 행이 없는 문체부 파일입니다: {path.name}")
        archive_rows[str(year)] = len(annual)
        records.extend(annual)
    frame = pd.DataFrame(records).sort_values(["plan_year", "province", "district", "title"]).reset_index(drop=True)
    return frame, {
        "source_url": SOURCE_URL, "years": sorted(seen_years), "archive_rows": archive_rows,
        "records": len(frame), "positive_attendance_rows": int((frame.prior_attendance.fillna(0) > 0).sum()),
        "measurement_counts": frame.measurement.value_counts(dropna=False).to_dict(),
    }


def feature_values(record: pd.Series | dict[str, Any]) -> dict[str, float]:
    prior = float(record["prior_attendance"])
    budget = float(record["budget_million_krw"])
    year = int(record["plan_year"])
    first_year = record.get("first_year")
    age = max(0, min(200, year - int(first_year))) if pd.notna(first_year) else 0
    month = int(record["month"])
    result = {
        "log_prior_attendance": float(np.log1p(prior)),
        "log_budget": float(np.log1p(budget)),
        "log_budget_per_prior_visitor": float(np.log1p(budget * 1_000_000 / prior)),
        "festival_age": float(age),
        "month_sin": float(np.sin(2 * np.pi * (month - 1) / 12)),
        "month_cos": float(np.cos(2 * np.pi * (month - 1) / 12)),
        "duration_days": float(record["duration_days"]),
    }
    province = record["province"] if record["province"] in PROVINCES else "unknown"
    festival_type = record["festival_type"] if record["festival_type"] in FESTIVAL_TYPES else "unknown"
    result.update({f"province_{index}": float(value == province) for index, value in enumerate(PROVINCES)})
    result.update({f"festival_type_{index}": float(value == festival_type) for index, value in enumerate(FESTIVAL_TYPES)})
    if list(result) != FEATURES or not np.isfinite(list(result.values())).all():
        raise ValueError("축제 관람수요 모델 입력을 만들지 못했습니다.")
    return result


def build_training_pairs(records: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, Any]]:
    """y년 계획과 y+1년 파일의 전년도 실적을 정확 일치시켜 y년 정답을 만든다."""
    rows: list[dict[str, Any]] = []
    excluded: Counter[str] = Counter()
    years = sorted(records.plan_year.unique())
    for year in years:
        following = records.loc[records.plan_year == year + 1]
        if following.empty:
            continue
        targets: dict[tuple[str, str, str], list[pd.Series]] = {}
        for _, target in following.iterrows():
            key = (target.province, target.district, target.title_key)
            targets.setdefault(key, []).append(target)
        for _, source in records.loc[records.plan_year == year].iterrows():
            key = (source.province, source.district, source.title_key)
            matches = targets.get(key, [])
            if len(matches) != 1:
                excluded["missing_or_ambiguous_adjacent_match"] += 1
                continue
            target = matches[0]
            if not pd.notna(source.prior_attendance) or not np.isfinite(source.prior_attendance) or source.prior_attendance <= 0:
                excluded["missing_nonpositive_prior_attendance"] += 1
                continue
            if not pd.notna(source.budget_million_krw) or not np.isfinite(source.budget_million_krw) or source.budget_million_krw < 0:
                excluded["missing_invalid_budget"] += 1
                continue
            if not pd.notna(target.prior_attendance) or not np.isfinite(target.prior_attendance) or target.prior_attendance <= 0:
                excluded["missing_nonpositive_target"] += 1
                continue
            features = feature_values(source)
            rows.append({
                "event_key": f"{year}:{source.province}:{source.district}:{source.title_key}",
                "target_year": year, "label_source_year": year + 1,
                "province": source.province, "district": source.district,
                "title": source.title, "title_key": source.title_key,
                "festival_type": source.festival_type,
                "prior_attendance": float(source.prior_attendance),
                "target_attendance": float(target.prior_attendance),
                "label_measurement": target.measurement,
                "budget_million_krw": float(source.budget_million_krw),
                "log_residual": float(np.log1p(target.prior_attendance) - np.log1p(source.prior_attendance)),
                **features,
            })
    if not rows:
        raise ValueError("인접 연도에 정확히 연결되는 축제 학습쌍이 없습니다.")
    frame = pd.DataFrame(rows).sort_values(["target_year", "event_key"]).reset_index(drop=True)
    return frame, {
        "feature_version": FEATURE_VERSION, "pairs": len(frame),
        "pairs_by_target_year": {str(k): int(v) for k, v in frame.groupby("target_year").size().items()},
        "excluded": dict(excluded), "matching": "동일 광역·기초지자체와 회차·연도를 제거한 축제명의 유일한 정확 일치",
        "limitations": [
            "지자체 제출 방문객 수로 계측·추정 방식이 섞여 있으며 2026 파일에서만 계측 방법을 제공했습니다.",
            "동일 축제의 전년도 실적이 있는 반복 개최 축제만 지원하며 신규 축제에는 적용하지 않습니다.",
            "2017~2021 원본 방문객 수는 천 명 단위, 2022년 이후는 명 단위로 변환했습니다.",
            "퍼지 매칭을 쓰지 않아 이름이나 관할이 크게 바뀐 축제는 학습에서 제외했습니다.",
        ],
    }
